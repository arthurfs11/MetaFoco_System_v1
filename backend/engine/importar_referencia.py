"""Importa as abas de referência (FERIADOS, CDI DIARIO, IPCA, TLP, UMSELIC)
e os parâmetros de cada contrato (abas individuais) para o PostgreSQL.

Uso:
    python -m engine.importar_referencia <caminho_planilha>
"""
import sys
import os
import openpyxl
from datetime import datetime, date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from database import SessionLocal
from models.models import Feriado, CdiDiario, Ipca, Tlp, Umselic, ContratoParam

HEADER_VARIANTES = {"padrao": 16, "deslocado": 17}


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _to_float(v):
    if isinstance(v, (int, float)):
        return float(v)
    return None


def detectar_aba_contrato(ws):
    """Verifica se a aba parece ser de contrato (tem VALOR FINANCIADO em H3 e tabela DATA)."""
    h3 = ws["H3"].value
    if not isinstance(h3, (int, float)):
        return False
    if detectar_layout(ws) is None:
        return False
    return True


def detectar_layout(ws):
    """Detecta a linha do cabeçalho da tabela de amortização e a variante."""
    for nome, lineno in HEADER_VARIANTES.items():
        v = ws[f"D{lineno}"].value
        if v in ("DATA", "DATA 1", "'DATA'", "=DATA"):
            return lineno
    # conversão: procurar 'DATA' na coluna A em linhas 16-20
    for r in range(16, 21):
        v = ws.cell(row=r, column=1).value
        if str(v).strip().upper().replace("'", "") in ("DATA",):
            return r
    return None


def importar_feriados(session, ws):
    existentes = {r for (r,) in session.query(Feriado.data).all()}
    vistos = set()
    for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        d = _to_date(row[0])
        if d is None or d in existentes or d in vistos:
            continue
        vistos.add(d)
        nome = row[2] if row[2] else ""
        session.add(Feriado(data=d, dia_semana=row[1], nome=str(nome)))


def importar_cdi_diario(session, ws):
    existentes = {r for (r,) in session.query(CdiDiario.data).all()}
    vistos = set()
    # dados começam na linha 41; col O=DATA(index14), N=fator_decimal(index13), R=SELIC(index17)
    for row in ws.iter_rows(min_row=41, max_col=19, values_only=True):
        d = _to_date(row[14])  # O=DATA
        if d is None or d in existentes or d in vistos:
            continue
        vistos.add(d)
        session.add(CdiDiario(
            data=d,
            fator_diario=_to_float(row[11]),
            fator_decimal=_to_float(row[13]),
            selic=_to_float(row[17]),
        ))


def importar_ipca(session, ws):
    existentes = {(m, a) for (m, a) in session.query(Ipca.mes, Ipca.ano).all()}
    vistos = set()
    for row in ws.iter_rows(min_row=2, max_col=6, values_only=True):
        mes = _to_float(row[1])
        ano = _to_float(row[2])
        if mes is None or ano is None:
            continue
        chave = (int(mes), int(ano))
        if chave in existentes or chave in vistos:
            continue
        vistos.add(chave)
        session.add(Ipca(
            mes=chave[0], ano=chave[1],
            indice=_to_float(row[3]), indice1=_to_float(row[4]),
        ))


def importar_tlp(session, ws):
    existentes = {r for (r,) in session.query(Tlp.data).all()}
    vistos = set()
    # dados começam na linha 2, col A=DATA, col D=TAXA(aa)
    for row in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        d = _to_date(row[0])
        if d is None or d in existentes or d in vistos:
            continue
        vistos.add(d)
        session.add(Tlp(
            data=d, mes=int(d.month), ano=int(d.year),
            taxa_aa=_to_float(row[3]),
        ))


def importar_umselic(session, ws):
    existentes = {r for (r,) in session.query(Umselic.data).all()}
    vistos = set()
    for row in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        d = _to_date(row[0])
        if d is None or d in existentes or d in vistos:
            continue
        vistos.add(d)
        session.add(Umselic(data=d, moeda=_to_float(row[3])))


def extrair_parametros_contrato(ws, nome_aba):
    """Extrai parâmetros das linhas 1-15 da aba de contrato."""
    def g(linha, col="H"):
        return ws[f"{col}{linha}"].value

    hdr = detectar_layout(ws)
    variante = "padrao"

    # detectar variante de conversão (colunas renomeadas)
    if hdr:
        p16 = ws.cell(row=hdr, column=16).value
        if p16 is not None and "CONV" in str(p16).upper():
            variante = "conversao"

    return ContratoParam(
        nome_aba=nome_aba,
        contrato=str(g(2)) if g(2) is not None else None,
        modalidade=str(g(1)) if g(1) is not None else None,
        valor_financiado=_to_float(g(3)),
        taxa_aa=_to_float(g(4)),
        indexador_nome=str(g(5, "E")) if g(5, "E") is not None else None,
        indexador_aa=_to_float(g(5)),
        sist_amort=str(g(5, "J")) if g(5, "J") is not None else None,
        usa_indice=str(g(5, "I")) if g(5, "I") is not None else None,
        taxa_am=_to_float(g(6)),
        taxa_ad=_to_float(g(8)),
        vencimento_final=_to_date(g(9)),
        entrada_recurso=_to_date(g(10)),
        data_1o_pagam=_to_date(g(11)),
        data_pgto_1o_parc=_to_date(g(12)),
        qt_parc=_to_float(g(13)),
        carencia=_to_float(g(14)),
        parc_anuais=str(g(13, "I")) if g(13, "I") is not None else None,
        cet_aa=_to_float(g(12, "J")),
        layout_header=hdr,
        layout_variante=variante,
    )


def main(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True, keep_vba=True)
    session = SessionLocal()

    importar_feriados(session, wb["FERIADOS"])
    importar_cdi_diario(session, wb["CDI DIARIO"])
    importar_ipca(session, wb["IPCA"])
    importar_tlp(session, wb["TLP"])
    importar_umselic(session, wb["UMSELIC"])
    session.commit()
    print("Referências importadas.")

    n_contrato = 0
    n_outras = 0
    for nome in wb.sheetnames:
        ws = wb[nome]
        try:
            if not detectar_aba_contrato(ws):
                n_outras += 1
                continue
        except Exception:
            continue
        if session.query(ContratoParam).filter(ContratoParam.nome_aba == nome).first():
            continue
        try:
            p = extrair_parametros_contrato(ws, nome)
            session.add(p)
            n_contrato += 1
        except Exception as e:
            print(f"  erro em {nome}: {e}")
    session.commit()
    print(f"Contratos extraídos: {n_contrato}")
    print(f"Abas não-contrato: {n_outras}")
    session.close()


if __name__ == "__main__":
    main(sys.argv[1])
